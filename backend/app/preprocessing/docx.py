"""DOCX / Word preprocessor adapter.

Conforms to Section 5.2 of ProjectDetails.md:
Extracts native XML structure (headings, tables, paragraphs) rather than blind OCR.
"""

from pathlib import Path
import uuid
from app.domain.document import Document, DocumentMetadata, DocumentType
from app.domain.provenance import Provenance
from app.domain.section import Block, BlockType, Section
from app.preprocessing.base import BasePreprocessor
from app.utils.hashing import compute_sha256_file


class DocxPreprocessor(BasePreprocessor):
    """Processes DOCX files into canonical documents via structural extraction."""

    def process(self, file_path: str | Path, document_id: str) -> Document:
        path = Path(file_path)
        source_hash = compute_sha256_file(path)
        file_size = path.stat().st_size

        metadata = DocumentMetadata(
            title=path.stem,
            source_path=str(path),
            file_type=DocumentType.DOCX,
            file_size_bytes=file_size,
            source_hash=source_hash,
        )

        doc_prov = Provenance(
            file_id=document_id,
            document_id=document_id,
            file_name=path.name,
            file_type="docx",
        )

        blocks: list[Block] = []
        md_lines: list[str] = [f"# {path.stem}\n"]

        try:
            import docx
            doc = docx.Document(str(path))

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                style_name = (p.style.name if p.style else "").lower()
                is_heading = "heading" in style_name or "title" in style_name
                btype = BlockType.HEADING if is_heading else BlockType.TEXT

                if is_heading:
                    md_lines.append(f"\n## {text}\n")
                else:
                    md_lines.append(f"{text}\n")

                blocks.append(
                    Block(
                        block_id=f"blk_{uuid.uuid4().hex[:8]}",
                        block_type=btype,
                        content=text,
                        provenance=doc_prov,
                    )
                )

            for table in doc.tables:
                rows_data = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    rows_data.append(" | ".join(row_cells))
                if rows_data:
                    table_str = "\n".join(rows_data)
                    md_lines.append(f"\n| {table_str} |\n")
                    blocks.append(
                        Block(
                            block_id=f"blk_{uuid.uuid4().hex[:8]}",
                            block_type=BlockType.TABLE,
                            content=table_str,
                            provenance=doc_prov,
                        )
                    )
        except Exception:
            pass

        if not blocks:
            # Fallback block
            fallback_text = f"DOCX document {path.name}"
            blocks.append(
                Block(
                    block_id=f"blk_{uuid.uuid4().hex[:8]}",
                    block_type=BlockType.TEXT,
                    content=fallback_text,
                    provenance=doc_prov,
                )
            )
            md_lines.append(fallback_text)

        section = Section(
            section_id=f"sec_{document_id}_main",
            title=path.stem,
            level=1,
            blocks=blocks,
            provenance=doc_prov,
        )

        normalized_md = "\n".join(md_lines)

        return Document(
            document_id=document_id,
            metadata=metadata,
            sections=[section],
            normalized_markdown=normalized_md,
            provenance=doc_prov,
        )
