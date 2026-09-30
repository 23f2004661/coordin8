"""High-fidelity multi-format artifact and document generation engine for Coordin8 projects.

Generates authentic, professionally styled documents in various stages of completion:
- Word Documents (.docx) with structured tables, callouts, and sign-offs
- Excel Workbooks (.xlsx) with multiple sheets, formulas, and KPI summaries
- PowerPoint Decks (.pptx) with modern slide layouts and metric cards
- PDF Whitepapers & Audit Dossiers (.pdf) with formal formatting
- Architecture Topology & Workflow Diagrams (.png) in crisp high-DPI
"""

from datetime import datetime
from pathlib import Path
from typing import Any
import os

# Third-party document libraries
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

import pptx
from pptx.dml.color import RGBColor as PPTX_RGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PPT_Inches, Pt as PPT_Pt

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import matplotlib
matplotlib.use("Agg")
import matplotlib.patches as patches
import matplotlib.pyplot as plt


def set_cell_background(cell, hex_color: str):
    """Set background color of a Word document table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding for a Word table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


# =========================================================================
# 1. WORD (.DOCX) GENERATOR
# =========================================================================
def generate_docx_document(
    output_path: Path,
    title: str,
    subtitle: str,
    project_name: str,
    client_name: str,
    stage_label: str,
    version: str,
    sections: list[dict[str, Any]],
    accent_hex: str = "4f46e5",
) -> Path:
    """Generate a formal multi-section Word document with executive styling."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()

    # Page Margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Accent RGB
    r = int(accent_hex[0:2], 16)
    g = int(accent_hex[2:4], 16)
    b = int(accent_hex[4:6], 16)
    accent_rgb = RGBColor(r, g, b)

    # Document Header / Pre-title
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run(f"COORDIN8 ENTERPRISE INTELLIGENCE  •  {project_name.upper()}  •  {stage_label.upper()}")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9)
    r_meta.font.bold = True
    r_meta.font.color.rgb = RGBColor(120, 130, 150)
    p_meta.paragraph_format.space_after = Pt(4)

    # Main Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = accent_rgb
    p_title.paragraph_format.space_after = Pt(2)

    # Subtitle
    if subtitle:
        p_sub = doc.add_paragraph()
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(13)
        r_sub.font.color.rgb = RGBColor(90, 100, 115)
        p_sub.paragraph_format.space_after = Pt(14)

    # Meta Table Box
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Client Organization", client_name), ("Document Version", version), ("Delivery Stage", stage_label)],
        [("Audited Date", datetime.now().strftime("%B %d, %Y")), ("Classification", "Confidential / Commercial"), ("Author", "Coordin8 Enterprise Team")],
    ]
    for row_idx, row_items in enumerate(meta_data):
        for col_idx, (label, val) in enumerate(row_items):
            cell = meta_table.cell(row_idx, col_idx)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            run_lbl = p.add_run(f"{label}\n")
            run_lbl.font.size = Pt(8)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = RGBColor(100, 116, 139)
            run_val = p.add_run(val)
            run_val.font.size = Pt(10)
            run_val.font.bold = True
            run_val.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Content Sections
    for sec in sections:
        sec_title = sec.get("title", "")
        if sec_title:
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            rh = h.add_run(sec_title)
            rh.font.name = "Calibri"
            rh.font.size = Pt(15)
            rh.font.bold = True
            rh.font.color.rgb = accent_rgb

        body = sec.get("body", "")
        if body:
            pb = doc.add_paragraph()
            pb.paragraph_format.space_after = Pt(8)
            pb.paragraph_format.line_spacing = 1.15
            rb = pb.add_run(body)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.font.color.rgb = RGBColor(51, 65, 85)

        bullets = sec.get("bullets", [])
        for b_text in bullets:
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_after = Pt(3)
            bp.paragraph_format.line_spacing = 1.15
            rb = bp.add_run(b_text)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10)
            rb.font.color.rgb = RGBColor(51, 65, 85)

        table_data = sec.get("table", None)
        if table_data and len(table_data) > 0:
            headers = table_data[0]
            rows = table_data[1:]
            t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            # Header Row
            for col_i, h_val in enumerate(headers):
                cell = t.cell(0, col_i)
                set_cell_background(cell, accent_hex)
                set_cell_margins(cell, 100, 100, 120, 120)
                cp = cell.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                cr = cp.add_run(str(h_val))
                cr.font.name = "Calibri"
                cr.font.size = Pt(9.5)
                cr.font.bold = True
                cr.font.color.rgb = RGBColor(255, 255, 255)
            # Body Rows
            for row_i, r_data in enumerate(rows):
                bg = "FFFFFF" if row_i % 2 == 0 else "F8FAFC"
                for col_i, cell_val in enumerate(r_data):
                    cell = t.cell(row_i + 1, col_i)
                    set_cell_background(cell, bg)
                    set_cell_margins(cell, 80, 80, 120, 120)
                    cp = cell.paragraphs[0]
                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    cr = cp.add_run(str(cell_val))
                    cr.font.name = "Calibri"
                    cr.font.size = Pt(9.5)
                    cr.font.color.rgb = RGBColor(30, 41, 59)
            doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Sign-off Approval Block
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_labels = [
        [("Client Executive Sign-off", f"{client_name} Authorized Representative"), ("Coordin8 Delivery Lead", "Srinath Srinivasan / Principal Architect")],
        [("Signature: __________________________", f"Status: {stage_label.upper()}"), ("Signature: __________________________", f"Date: {datetime.now().strftime('%Y-%m-%d')}")],
    ]
    for r_i, r_list in enumerate(sign_labels):
        for c_i, (t1, t2) in enumerate(r_list):
            c = sign_table.cell(r_i, c_i)
            set_cell_background(c, "F1F5F9")
            set_cell_margins(c, 100, 100, 120, 120)
            p = c.paragraphs[0]
            r1 = p.add_run(f"{t1}\n")
            r1.font.bold = True
            r1.font.size = Pt(9)
            r2 = p.add_run(t2)
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(str(output_path))
    return output_path


# =========================================================================
# 2. EXCEL (.XLSX) GENERATOR
# =========================================================================
def generate_xlsx_workbook(
    output_path: Path,
    title: str,
    project_name: str,
    client_name: str,
    sheets_data: list[dict[str, Any]],
    accent_hex: str = "4F46E5",
) -> Path:
    """Generate a multi-tab formatted Excel workbook with formulas, KPIs, and styled tables."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_fill = PatternFill(start_color=accent_hex, end_color=accent_hex, fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    kpi_title_font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    kpi_val_font = Font(name="Segoe UI", size=18, bold=True, color="0F172A")
    data_font = Font(name="Segoe UI", size=9.5, color="1E293B")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    kpi_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for sheet_spec in sheets_data:
        sheet_name = sheet_spec.get("name", "Sheet1")[:30]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        # Sheet Title Banner
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"{title.upper()}  •  {sheet_name.upper()}"
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1E293B")
        title_cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 32

        # Subtitle Row
        ws.merge_cells("A2:G2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Client: {client_name}  |  Project: {project_name}  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        sub_cell.font = Font(name="Segoe UI", size=9.5, color="64748B")
        sub_cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[2].height = 20

        current_row = 4

        # Optional KPI Summary Cards
        kpis = sheet_spec.get("kpis", [])
        if kpis:
            col_start = 1
            for k in kpis:
                c_lbl = ws.cell(row=current_row, column=col_start, value=k.get("label", ""))
                c_lbl.font = kpi_title_font
                c_lbl.fill = kpi_fill
                c_lbl.alignment = Alignment(horizontal="center", vertical="center")
                c_lbl.border = thin_border

                c_val = ws.cell(row=current_row + 1, column=col_start, value=k.get("value", ""))
                c_val.font = kpi_val_font
                c_val.fill = kpi_fill
                c_val.alignment = Alignment(horizontal="center", vertical="center")
                c_val.border = thin_border

                ws.column_dimensions[get_column_letter(col_start)].width = 20
                col_start += 1

            ws.row_dimensions[current_row].height = 20
            ws.row_dimensions[current_row + 1].height = 36
            current_row += 3

        # Main Table
        table_headers = sheet_spec.get("headers", [])
        table_rows = sheet_spec.get("rows", [])

        if table_headers:
            for col_idx, h_text in enumerate(table_headers, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=h_text)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                c.border = thin_border
            ws.row_dimensions[current_row].height = 26
            current_row += 1

            for r_idx, r_data in enumerate(table_rows):
                is_alt = (r_idx % 2 == 1)
                for c_idx, val in enumerate(r_data, start=1):
                    c = ws.cell(row=current_row, column=c_idx, value=val)
                    c.font = data_font
                    c.border = thin_border
                    if is_alt:
                        c.fill = alt_fill
                    # Number format detection
                    if isinstance(val, (int, float)):
                        if str(table_headers[c_idx - 1]).lower().endswith(("%", "rate", "pct", "ratio")):
                            c.number_format = "0.0%"
                        elif str(table_headers[c_idx - 1]).lower().startswith(("$", "revenue", "cost", "loan", "balance", "val")):
                            c.number_format = "$#,##0.00"
                        else:
                            c.number_format = "#,##0"
                    elif str(val).startswith("="):
                        c.number_format = "$#,##0.00"

                ws.row_dimensions[current_row].height = 20
                current_row += 1

            # Summary Formula Row if numeric columns exist
            if sheet_spec.get("include_totals", False) and len(table_rows) > 0:
                summary_row = current_row
                ws.cell(row=summary_row, column=1, value="TOTAL / SUMMARY").font = Font(name="Segoe UI", size=10, bold=True)
                ws.cell(row=summary_row, column=1).border = thin_border
                for c_idx in range(2, len(table_headers) + 1):
                    col_letter = get_column_letter(c_idx)
                    start_r = summary_row - len(table_rows)
                    end_r = summary_row - 1
                    # Check if column is numeric
                    first_val = table_rows[0][c_idx - 1]
                    if isinstance(first_val, (int, float)):
                        c = ws.cell(row=summary_row, column=c_idx, value=f"=SUM({col_letter}{start_r}:{col_letter}{end_r})")
                        c.font = Font(name="Segoe UI", size=10, bold=True)
                        c.border = thin_border
                        c.fill = alt_fill
                        if str(table_headers[c_idx - 1]).lower().startswith(("$", "rev", "cost", "loan", "val")):
                            c.number_format = "$#,##0.00"
                        else:
                            c.number_format = "#,##0"
                    else:
                        ws.cell(row=summary_row, column=c_idx).border = thin_border
                ws.row_dimensions[summary_row].height = 24

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and not cell.coordinate in ["A1", "A2"]:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(str(output_path))
    return output_path


# =========================================================================
# 3. POWERPOINT (.PPTX) GENERATOR
# =========================================================================
def generate_pptx_deck(
    output_path: Path,
    title: str,
    subtitle: str,
    project_name: str,
    client_name: str,
    stage_label: str,
    slides_data: list[dict[str, Any]],
    accent_rgb_tuple: tuple[int, int, int] = (79, 70, 229),
) -> Path:
    """Generate a sleek executive presentation deck with metric cards and structured takeaways."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = pptx.Presentation()
    prs.slide_width = PPT_Inches(13.333)
    prs.slide_height = PPT_Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    accent_color = PPTX_RGB(*accent_rgb_tuple)
    dark_bg = PPTX_RGB(15, 23, 42)
    white = PPTX_RGB(255, 255, 255)
    light_text = PPTX_RGB(148, 163, 184)
    dark_text = PPTX_RGB(30, 41, 59)

    # 1. Title Slide (Dark Theme)
    title_slide = prs.slides.add_slide(blank_layout)
    bg_shape = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = dark_bg
    bg_shape.line.fill.background()

    # Pre-title Tag
    tx = title_slide.shapes.add_textbox(PPT_Inches(1.2), PPT_Inches(1.8), PPT_Inches(10.5), PPT_Inches(0.5))
    tf = tx.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = f"COORDIN8 EXECUTIVE STRATEGY  •  {project_name.upper()}  •  {stage_label.upper()}"
    p0.font.size = PPT_Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = accent_color

    # Main Title
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.size = PPT_Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = white
    p1.space_before = PPT_Pt(10)

    # Subtitle
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = PPT_Pt(18)
        p2.font.color.rgb = light_text
        p2.space_before = PPT_Pt(12)

    # Meta Footer on Title Slide
    p3 = tf.add_paragraph()
    p3.text = f"Client: {client_name}   |   Date: {datetime.now().strftime('%B %d, %Y')}   |   Status: {stage_label}"
    p3.font.size = PPT_Pt(12)
    p3.font.color.rgb = light_text
    p3.space_before = PPT_Pt(36)

    # 2. Content Slides
    for slide_spec in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        # Header text
        header_box = slide.shapes.add_textbox(PPT_Inches(1.0), PPT_Inches(0.6), PPT_Inches(11.3), PPT_Inches(1.0))
        htf = header_box.text_frame
        hp = htf.paragraphs[0]
        hp.text = slide_spec.get("title", "Slide Title")
        hp.font.size = PPT_Pt(24)
        hp.font.bold = True
        hp.font.color.rgb = dark_text

        hsub = htf.add_paragraph()
        hsub.text = slide_spec.get("category", project_name.upper())
        hsub.font.size = PPT_Pt(10)
        hsub.font.bold = True
        hsub.font.color.rgb = accent_color
        hsub.space_before = PPT_Pt(3)

        # Content Cards or Points
        cards = slide_spec.get("cards", [])
        if cards:
            num_cards = min(len(cards), 3)
            card_w = (11.333 - (num_cards - 1) * 0.4) / num_cards
            top_y = 1.8
            card_h = 4.8

            for c_i, card in enumerate(cards[:3]):
                left_x = 1.0 + c_i * (card_w + 0.4)
                # Card box shape
                card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PPT_Inches(left_x), PPT_Inches(top_y), PPT_Inches(card_w), PPT_Inches(card_h))
                card_shape.fill.solid()
                card_shape.fill.fore_color.rgb = PPTX_RGB(248, 250, 252)
                card_shape.line.color.rgb = PPTX_RGB(226, 232, 240)

                ctf = card_shape.text_frame
                ctf.word_wrap = True
                ctf.margin_left = PPT_Inches(0.3)
                ctf.margin_right = PPT_Inches(0.3)
                ctf.margin_top = PPT_Inches(0.3)

                cp0 = ctf.paragraphs[0]
                cp0.text = card.get("title", "")
                cp0.font.size = PPT_Pt(16)
                cp0.font.bold = True
                cp0.font.color.rgb = accent_color

                if "metric" in card:
                    cp_m = ctf.add_paragraph()
                    cp_m.text = str(card["metric"])
                    cp_m.font.size = PPT_Pt(30)
                    cp_m.font.bold = True
                    cp_m.font.color.rgb = dark_text
                    cp_m.space_before = PPT_Pt(8)

                cp1 = ctf.add_paragraph()
                cp1.text = card.get("desc", "")
                cp1.font.size = PPT_Pt(11)
                cp1.font.color.rgb = PPTX_RGB(71, 85, 105)
                cp1.space_before = PPT_Pt(10)

                for b in card.get("bullets", []):
                    bp = ctf.add_paragraph()
                    bp.text = f"• {b}"
                    bp.font.size = PPT_Pt(10.5)
                    bp.font.color.rgb = PPTX_RGB(51, 65, 85)
                    bp.space_before = PPT_Pt(4)

        # Slide Footer
        footer_box = slide.shapes.add_textbox(PPT_Inches(1.0), PPT_Inches(6.8), PPT_Inches(11.3), PPT_Inches(0.4))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = f"Coordin8 Delivery System  •  Confidential for {client_name}  •  {stage_label}"
        fp.font.size = PPT_Pt(9)
        fp.font.color.rgb = light_text

    prs.save(str(output_path))
    return output_path


# =========================================================================
# 4. PDF GENERATOR
# =========================================================================
def generate_pdf_document(
    output_path: Path,
    title: str,
    project_name: str,
    client_name: str,
    stage_label: str,
    sections: list[dict[str, Any]],
    accent_hex: str = "#4F46E5",
) -> Path:
    """Generate a formal PDF whitepaper and compliance report with ReportLab."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()
    accent_c = colors.HexColor(accent_hex)
    dark_c = colors.HexColor("#0F172A")
    body_c = colors.HexColor("#334155")
    meta_c = colors.HexColor("#64748B")

    title_style = ParagraphStyle(
        "PDFTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=accent_c,
        spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "PDFMeta",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=12,
        textColor=meta_c,
        spaceAfter=12,
    )
    h2_style = ParagraphStyle(
        "PDFH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=dark_c,
        spaceBefore=14,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "PDFBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        textColor=body_c,
        spaceAfter=8,
    )
    bullet_style = ParagraphStyle(
        "PDFBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=body_c,
        leftIndent=14,
        spaceAfter=4,
    )

    story = []

    # Header pre-title
    story.append(Paragraph(f"COORDIN8 VERIFIED ARTIFACT  •  {project_name.upper()}  •  {stage_label.upper()}", meta_style))
    story.append(Paragraph(title, title_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=accent_c, spaceAfter=14))

    # Meta Overview Box
    meta_table_data = [
        [
            Paragraph(f"<b>Client:</b> {client_name}", body_style),
            Paragraph(f"<b>Audit Date:</b> {datetime.now().strftime('%B %d, %Y')}", body_style),
            Paragraph(f"<b>Status:</b> {stage_label}", body_style),
        ]
    ]
    meta_table = Table(meta_table_data, colWidths=[180, 160, 160])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 14))

    for sec in sections:
        sec_title = sec.get("title", "")
        if sec_title:
            story.append(Paragraph(sec_title, h2_style))
        body = sec.get("body", "")
        if body:
            story.append(Paragraph(body, body_style))
        for b_text in sec.get("bullets", []):
            story.append(Paragraph(f"• {b_text}", bullet_style))

        table_data = sec.get("table", None)
        if table_data and len(table_data) > 0:
            formatted_table = []
            for row_i, r in enumerate(table_data):
                row_cells = []
                for cell_val in r:
                    style_use = ParagraphStyle(
                        f"cell_{row_i}",
                        parent=body_style,
                        fontSize=8.5,
                        leading=11,
                        textColor=colors.white if row_i == 0 else body_c,
                        fontName="Helvetica-Bold" if row_i == 0 else "Helvetica",
                    )
                    row_cells.append(Paragraph(str(cell_val), style_use))
                formatted_table.append(row_cells)

            col_w = 504 / len(table_data[0])
            t = Table(formatted_table, colWidths=[col_w] * len(table_data[0]))
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), accent_c),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

    # Formal Sign-off Box
    story.append(Spacer(1, 16))
    sign_data = [
        [
            Paragraph("<b>Executive Sign-off & Client Acceptance:</b><br/>"
                      f"Authorized Signature for {client_name}<br/><br/>"
                      "________________________________________<br/>"
                      f"Date: {datetime.now().strftime('%Y-%m-%d')}", body_style),
            Paragraph("<b>Coordin8 Delivery & QA Certification:</b><br/>"
                      "Srinath Srinivasan, Lead AI Architect<br/><br/>"
                      "________________________________________<br/>"
                      f"Status: {stage_label.upper()}", body_style),
        ]
    ]
    sign_table = Table(sign_data, colWidths=[250, 254])
    sign_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
        ("PADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(sign_table)

    doc.build(story)
    return output_path


# =========================================================================
# 5. DIAGRAM (.PNG) GENERATOR
# =========================================================================
def generate_architecture_diagram(
    output_path: Path,
    title: str,
    project_code: str,
    nodes: list[dict[str, Any]],
    theme_hex: str = "#4F46E5",
) -> Path:
    """Generate crisp high-DPI architecture workflow and system topology diagram."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(12, 6.5), dpi=220)
    fig.patch.set_facecolor("#0F172A")
    ax.set_facecolor("#0F172A")

    # Title & Subtitle Banner
    ax.text(0.05, 0.93, f"{title.upper()}", fontsize=15, fontweight="bold", color="#FFFFFF", transform=ax.transAxes)
    ax.text(0.05, 0.88, f"Coordin8 Systems Architecture  •  {project_code} Pipeline  •  High-Throughput Topology", fontsize=9.5, color="#94A3B8", transform=ax.transAxes)

    num_nodes = len(nodes)
    if num_nodes == 0:
        nodes = [
            {"title": "Ingestion", "desc": "Raw Streams"},
            {"title": "Processing", "desc": "Feature Extraction"},
            {"title": "Model Core", "desc": "Inference Engine"},
            {"title": "Output / API", "desc": "Decision Engine"},
        ]
        num_nodes = 4

    box_w = 0.18
    box_h = 0.45
    spacing = (0.90 - box_w) / max(num_nodes - 1, 1)

    for idx, node in enumerate(nodes):
        x = 0.05 + idx * spacing
        y = 0.25

        # Draw box
        rect = patches.FancyBboxPatch(
            (x, y),
            box_w,
            box_h,
            boxstyle="round,pad=0.03,rounding_size=0.02",
            facecolor="#1E293B",
            edgecolor=theme_hex,
            linewidth=2.0,
            transform=ax.transAxes,
        )
        ax.add_patch(rect)

        # Stage Number Badge
        ax.text(
            x + 0.02,
            y + box_h - 0.06,
            f"0{idx+1}",
            fontsize=8.5,
            fontweight="bold",
            color=theme_hex,
            transform=ax.transAxes,
        )

        # Node Title
        ax.text(
            x + 0.02,
            y + box_h - 0.14,
            node.get("title", f"Stage {idx+1}"),
            fontsize=11,
            fontweight="bold",
            color="#F8FAFC",
            transform=ax.transAxes,
        )

        # Node Description
        ax.text(
            x + 0.02,
            y + 0.08,
            node.get("desc", ""),
            fontsize=8.5,
            color="#94A3B8",
            wrap=True,
            transform=ax.transAxes,
        )

        # Draw connecting arrow to next node
        if idx < num_nodes - 1:
            ax.annotate(
                "",
                xy=(x + box_w + spacing * 0.42, y + box_h / 2),
                xytext=(x + box_w + 0.02, y + box_h / 2),
                xycoords="axes fraction",
                arrowprops=dict(arrowstyle="->", color=theme_hex, lw=2.5, mutation_scale=15),
            )

    # Status Bar at bottom
    ax.text(0.05, 0.06, "Architecture Status: VALIDATED & BENCHMARKED  |  SLA: <400ms P99  |  HA Active-Active Multi-AZ", fontsize=8.5, color="#64748B", transform=ax.transAxes)

    ax.axis("off")
    plt.tight_layout()
    plt.savefig(str(output_path), facecolor=fig.get_facecolor(), edgecolor="none")
    plt.close(fig)
    return output_path
